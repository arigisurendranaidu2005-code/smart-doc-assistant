"""Document ingestion and chunking module."""
import logging
import os
from typing import List, Dict, Any, Optional
from pathlib import Path

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

from src.config import settings

logger = logging.getLogger(__name__)

class DocumentIngestor:
    """Handles loading and chunking of documents."""

    def __init__(self, chunk_size: int = None, chunk_overlap: int = None):
        """Initialize with chunking parameters."""
        self.chunk_size = chunk_size or settings.chunk_size
        self.chunk_overlap = chunk_overlap or settings.chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            separators=["\n\n", "\n", ".", " ", ""]
        )

    def load_document(self, file_path: str | Path) -> List[Document]:
        """Load a document and return a list of parsed Documents."""
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        logger.info(f"Loading document: {file_path}")
        ext = file_path.suffix.lower()

        if ext == '.pdf':
            return self._load_pdf(file_path)
        elif ext == '.docx':
            return self._load_docx(file_path)
        elif ext == '.txt':
            return self._load_txt(file_path)
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

    def _load_pdf(self, file_path: Path) -> List[Document]:
        """Load PDF file using PyPDF2."""
        if PdfReader is None:
            raise ImportError("PyPDF2 is required to process PDF files. Run: pip install PyPDF2")
        
        documents = []
        try:
            with open(file_path, "rb") as f:
                reader = PdfReader(f)
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        metadata = {
                            "source": str(file_path),
                            "filename": file_path.name,
                            "page_number": i + 1,
                        }
                        documents.append(Document(page_content=text, metadata=metadata))
        except Exception as e:
            logger.error(f"Error loading PDF {file_path}: {e}")
            raise
        return documents

    def _load_docx(self, file_path: Path) -> List[Document]:
        """Load DOCX file using python-docx."""
        if docx is None:
            raise ImportError("python-docx is required to process DOCX files. Run: pip install python-docx")
        
        try:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            text = "\n".join(full_text)
            
            metadata = {
                "source": str(file_path),
                "filename": file_path.name,
            }
            return [Document(page_content=text, metadata=metadata)]
        except Exception as e:
            logger.error(f"Error loading DOCX {file_path}: {e}")
            raise

    def _load_txt(self, file_path: Path) -> List[Document]:
        """Load TXT file."""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
            metadata = {
                "source": str(file_path),
                "filename": file_path.name,
            }
            return [Document(page_content=text, metadata=metadata)]
        except Exception as e:
            logger.error(f"Error loading TXT {file_path}: {e}")
            raise

    def process_document(self, file_path: str | Path) -> List[Document]:
        """Load and chunk a document."""
        docs = self.load_document(file_path)
        logger.info(f"Chunking {len(docs)} pages/sections from {file_path}")
        
        chunked_docs = self.text_splitter.split_documents(docs)
        
        # Add chunk index to metadata
        for i, doc in enumerate(chunked_docs):
            doc.metadata["chunk_index"] = i
            
        logger.info(f"Created {len(chunked_docs)} chunks from {file_path}")
        return chunked_docs
